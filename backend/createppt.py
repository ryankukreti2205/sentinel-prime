from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# The Exact 20-Slide Dictionary
slides_data = [
    # PART 1: INTRO TO DL (1-7)
    {"title": "Sentinel Prime: Final Project", "bullets": ["AI-Powered Emergency Vehicle Preemption", "Presenter: [Your Name]"]},
    {"title": "Basics of Deep Learning & ANNs", "bullets": ["Deep Learning uses multi-layered neural networks to extract features.", "Architecture: Input Layer -> Hidden Layers -> Output Layer."]},
    {"title": "Forward Propagation & Activations", "bullets": ["Forward Pass: Calculates weighted sums using Weights and Biases.", "Activation Functions (ReLU, Sigmoid) introduce non-linearity."]},
    {"title": "Learning via Backpropagation", "bullets": ["Loss Functions (MSE, Cross-Entropy) measure the network's error.", "Optimizers (Adam) use Gradient Descent to adjust weights backward."]},
    {"title": "Convolutional Neural Networks (CNN)", "bullets": ["Standard ANNs flatten images and lose spatial layout.", "CNNs use sliding Kernels (filters) to detect edges and shapes.", "Pooling layers reduce computational size."]},
    {"title": "RNNs & Object Detection", "bullets": ["RNNs handle sequential data like video frames.", "YOLO (You Only Look Once) revolutionized computer vision by processing whole images in a single fast pass."]},
    {"title": "DL Applications in Smart Cities", "bullets": ["Automated license plate recognition and pedestrian tracking.", "Our App: Real-time emergency vehicle identification to clear intersections."]},
    
    # PART 2: PROJECT WORK (8-19)
    {"title": "Problem Statement", "bullets": ["Intersections cause critical delays for emergency responders.", "Traditional acoustic and strobe sensors are flawed and expensive.", "Goal: A scalable, vision-based CCTV system to preempt traffic lights."]},
    {"title": "Literature Review", "bullets": ["Past Methods: RF emitters cost $2,000+ per vehicle to install.", "Our Approach: Edge-AI shifts intelligence to the camera.", "Requires zero additional hardware on the emergency vehicles."]},
    {"title": "Dataset Used", "bullets": ["Gathered ~4,500 images of ambulances, fire trucks, and police cars.", "Fixed 'Ghost Class IDs' and purged corrupted files.", "Applied Mosaic augmentation for robust feature learning.", "[ 📸 INSERT DATASET COLLAGE SCREENSHOT ]"]},
    {"title": "Methodology", "bullets": ["1. Edge Detection: YOLOv8 analyzes video locally.", "2. Logic Gate: Filters out visual noise.", "3. API Bridge: Threaded HTTP POST to backend.", "4. Command Center: React UI visually alerts operators."]},
    {"title": "Model Architecture & System Design", "bullets": ["Python Edge Node -> REST API -> MongoDB -> React Dashboard.", "Strict separation of heavy AI processing and lightweight web backend.", "[ 📸 INSERT ARCHITECTURE DIAGRAM ]"]},
    {"title": "Tools & Technologies Used", "bullets": ["AI/Vision: Python, YOLOv8, OpenCV, ByteTrack.", "Backend: Node.js, Express.js (REST API).", "Database: MongoDB Atlas (Cloud NoSQL).", "Frontend: React, Vite, Axios, Custom CSS."]},
    {"title": "Work Completed So Far", "bullets": ["Curated dataset and trained custom YOLOv8 model.", "Engineered the Anti-Flicker persistence logic script.", "Deployed the Node.js + MongoDB backend.", "Developed the live React video-streaming UI."]},
    {"title": "Results & Performance", "bullets": ["Achieved strong mAP50 score indicating reliable recognition.", "Zero False-Positives due to the strict Time-Gate requirement.", "[ 📸 INSERT YOLO RESULTS.PNG GRAPH ]"]},
    {"title": "Deep Dive: Anti-Flicker Logic", "bullets": ["Raw AI models 'flicker' on similar objects (like delivery vans).", "Solution: Target must be detected with >65% confidence for 15 frames.", "One missed frame hard-resets the counter to zero.", "[ 📸 INSERT TERMINAL LOG SCREENSHOT ]"]},
    {"title": "The Command Center (Frontend UI)", "bullets": ["React dashboard dynamically fetches new data every second.", "Users upload local traffic footage directly into the browser.", "Flask microservice streams AI-processed video back in real-time.", "[ 📸 INSERT REACT DASHBOARD SCREENSHOT ]"]},
    {"title": "Challenges Faced", "bullets": ["Corrupted Datasets: Debugging indexing errors during training.", "Hardware Constraints: Mapping PyTorch to Apple Silicon MPS GPU.", "Async Blocking: Overcame video freezes using Python threading."]},
    {"title": "Next Plan of Action", "bullets": ["Port the Python edge-script to a Jetson Nano or Raspberry Pi.", "Integrate physical GPIO pins to control a miniature hardware traffic light."]},
    
    # PART 3: GANTT CHART (20)
    {"title": "Project Timeline (Gantt Chart)", "bullets": ["Completed: Dataset, Training, Backend Setup, Full-Stack UI.", "Upcoming: Hardware Integration, Final Testing, Report Writing.", "[ 📸 INSERT GANTT CHART SCREENSHOT ]"]}
]

# Create Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
slide.shapes.title.text = slides_data[0]["title"]
slide.placeholders[1].text = "\n".join(slides_data[0]["bullets"])

# Create Content Slides
bullet_slide_layout = prs.slide_layouts[1]
for slide_info in slides_data[1:]:
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = slide_info["title"]
    tf = slide.shapes.placeholders[1].text_frame
    for i, bullet in enumerate(slide_info["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet

prs.save("Sentinel_Final_Presentation_20_Slides.pptx")
print("✅ Generated exact 20-slide presentation: Sentinel_Final_Presentation_20_Slides.pptx")